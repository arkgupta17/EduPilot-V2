from pathlib import Path

from docx import Document as DocxDocument

from app.models.document import Document


class DOCXProcessor:

    def process(self, file_path: str) -> Document:

        path = Path(file_path)

        doc = DocxDocument(file_path)

        sections = []
        current_heading = path.stem
        current_text = []

        def flush_section():
            if not current_text:
                return

            text = "\n".join(current_text).strip()

            if text:
                sections.append(
                    {
                        "heading": current_heading,
                        "text": text
                    }
                )

            current_text.clear()

        for paragraph in doc.paragraphs:

            text = paragraph.text.strip()

            if not text:
                continue

            style_name = (
                paragraph.style.name.lower()
                if paragraph.style
                else ""
            )

            if "heading" in style_name:

                flush_section()

                current_heading = text

            else:

                current_text.append(text)

        flush_section()

        # Extract tables
        for table_index, table in enumerate(
            doc.tables,
            start=1
        ):

            rows = []

            for row in table.rows:

                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                rows.append(
                    " | ".join(cells)
                )

            table_text = "\n".join(rows).strip()

            if table_text:

                sections.append(
                    {
                        "heading": f"Table {table_index}",
                        "text": table_text
                    }
                )

        content_parts = [
            section["text"]
            for section in sections
        ]

        content = "\n\n".join(content_parts)

        return Document(
            id=str(path.resolve()),
            title=path.stem,
            source_type="docx",
            url=str(path.resolve()),
            content=content,
            metadata={
                "source": str(path.resolve()),
                "filename": path.name,
                "sections": sections
            }
        )